"""Convert Cupboard Manager 3D version 11 exports to Cut Manager workbooks."""

from copy import copy
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re

from openpyxl import load_workbook


SOURCE_HEADERS = (
    "Unit ID", "L", "W", "Qty", "Board Category", "Board Name",
    "Edge L", "Edge W", "Hole L", "Hole W", "Edging Category",
    "Edging Name", "Remarks",
)


@dataclass
class Cutlist:
    board_category: str
    board_name: str
    thickness: int | float
    edging_category: str | None
    edging_name: str | None
    pieces: list[tuple]


def _text(value):
    return "" if value is None else str(value).strip()


def _info_sheet(workbook):
    return next((s for s in workbook if s.title.strip().casefold() == "cutlist info"), None)


def _has_signature(workbook):
    info = _info_sheet(workbook)
    if info is not None and _text(info["A2"].value).casefold().startswith("cupboard manager"):
        return True
    return any(_text(s["A1"].value).casefold().startswith("cupboard manager") for s in workbook)


def is_cupboard_cutlist(file_path):
    """Identify the format by its contents, independently of its version or filename."""
    if Path(file_path).suffix.lower() != ".xlsx":
        return False
    workbook = load_workbook(file_path, read_only=True)
    try:
        return _has_signature(workbook)
    finally:
        workbook.close()


def _error(sheet, row, message):
    return ValueError(f"{sheet.title}, row {row}: {message}")


def _number(value, sheet, row, field, *, integer=False, minimum=0, maximum=None):
    try:
        if isinstance(value, bool) or _text(value) == "":
            raise InvalidOperation
        number = Decimal(_text(value).replace(",", "."))
        if not number.is_finite() or number < minimum:
            raise InvalidOperation
        if integer and number != number.to_integral_value():
            raise InvalidOperation
        if maximum is not None and number > maximum:
            raise InvalidOperation
        result = int(number) if number == number.to_integral_value() else float(number)
        if not isinstance(result, int) and not Decimal(str(result)).is_finite():
            raise InvalidOperation
        return result
    except (InvalidOperation, ValueError, OverflowError):
        raise _error(sheet, row, f"invalid {field}: {value!r}.") from None


def _read_cutlist(sheet, start_row):
    for column, expected in enumerate(SOURCE_HEADERS, 1):
        if _text(sheet.cell(start_row - 1, column).value).casefold() != expected.casefold():
            raise _error(sheet, start_row - 1, f"expected column {column} heading {expected!r}.")
    category = _text(sheet["C2"].value)
    name = _text(sheet["E2"].value)
    if not category or not name:
        raise _error(sheet, 2, "board category and board name are required.")
    thickness = _number(sheet["K2"].value, sheet, 2, "board thickness", minimum=Decimal("0.000001"))
    pieces = []
    ended = False
    for row_number, row in enumerate(sheet.iter_rows(min_row=start_row, max_col=max(21, sheet.max_column)), start_row):
        values = [cell.value for cell in row]
        if not any(_text(value) for value in values):
            continue
        if ended:
            raise _error(sheet, row_number, "unexpected data after EndOfFile.")
        if _text(values[0]).casefold() == "endoffile":
            if any(_text(value) for value in values[1:]):
                raise _error(sheet, row_number, "unexpected data on EndOfFile row.")
            ended = True
            continue
        if any(_text(value) for value in values[13:]):
            raise _error(sheet, row_number, "groove, mitre, cutout or crosscut instructions are unsupported.")
        if _text(values[4]) != category or _text(values[5]) != name:
            raise _error(sheet, row_number, "piece board category/name differs from the sheet heading.")
        piece = [
            _number(values[0], sheet, row_number, "unit ID", integer=True),
            _number(values[1], sheet, row_number, "length", minimum=Decimal("0.000001")),
            _number(values[2], sheet, row_number, "width", minimum=Decimal("0.000001")),
            _number(values[3], sheet, row_number, "quantity", integer=True, minimum=1),
        ]
        for index, field in [(6, "Edge L"), (7, "Edge W"), (8, "Hole L"), (9, "Hole W")]:
            piece.append(_number(values[index], sheet, row_number, field, integer=True,
                                 maximum=2 if index in (6, 7) else None))
        for value in values[10:12]:
            piece.append(value if _text(value) else None)
        pieces.append(tuple(piece))
    if not pieces:
        raise _error(sheet, start_row, "cutlist contains no pieces.")
    return Cutlist(category, name, thickness, sheet["C3"].value, sheet["E3"].value, pieces)


def _read_input(file_path):
    workbook = load_workbook(file_path, read_only=True)
    try:
        if not _has_signature(workbook):
            raise ValueError("This is not a Cupboard Manager workbook.")
        info = _info_sheet(workbook)
        if info is None:
            raise ValueError("Cupboard Manager workbook is missing its Cutlist Info sheet.")
        version = info["B11"].value
        if _text(version) != "11":
            raise _error(info, 11, f"unsupported file version {version!r}; expected version 11.")
        start_row = _number(info["D20"].value, info, 20, "cutlist start row", integer=True, minimum=5)
        if start_row > 1048576:
            raise _error(info, 20, "cutlist start row is outside the worksheet.")
        filename = re.fullmatch(r"(?:\d+\s*-\s*)?(.*?)_Job-(.+?)_Plan-.+", Path(file_path).stem, re.IGNORECASE)
        customer = _text(info["D13"].value) or (filename[1].strip() if filename else "")
        job = filename[2].strip() if filename else ""
        if not customer or not job:
            raise ValueError("Cannot determine customer/job. Expected filename: <customer>_Job-<job>_Plan-<plan>.xlsx")
        cutlists = []
        for sheet in workbook:
            if sheet is info:
                continue
            if not re.fullmatch(r"CutList\d+", sheet.title, re.IGNORECASE):
                raise ValueError(f"Unsupported Cupboard Manager sheet: {sheet.title!r}.")
            cutlists.append(_read_cutlist(sheet, start_row))
        if not cutlists:
            raise ValueError("Cupboard Manager workbook contains no cutlists.")
        return customer, job, info["D14"].value, cutlists
    finally:
        workbook.close()


def _write_text(cell, value):
    # Names beginning with '=' are literal source names, never Excel formulas.
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def process_cupboard_cutlist(file_path, template_path):
    """Validate and convert a source workbook, returning its new Path.

    The app enforces machine licensing and developer access before calling this.
    All source sheets are validated before any output file is created.
    """
    customer, job, contact, cutlists = _read_input(file_path)
    workbook = load_workbook(template_path)
    try:
        while len(workbook.worksheets) < len(cutlists):
            workbook.copy_worksheet(workbook.worksheets[0])
        for sheet in list(workbook.worksheets[len(cutlists):]):
            workbook.remove(sheet)
        for index, cutlist in enumerate(cutlists):
            sheet = workbook.worksheets[index]
            sheet.title = f"CutList{index + 1}"
            for address, value in {
                "C2": customer, "C3": job, "J2": contact,
                "C4": cutlist.board_category, "E4": cutlist.board_name,
                "J4": cutlist.thickness, "C5": cutlist.edging_category,
                "E5": cutlist.edging_name,
            }.items():
                _write_text(sheet[address], value)
            styles = [copy(sheet.cell(8, col)._style) for col in range(1, 12)]
            row_height = sheet.row_dimensions[8].height
            for row in sheet.iter_rows(min_row=8):
                for cell in row:
                    cell.value = None
            for row_number, piece in enumerate(cutlist.pieces, 8):
                for column, value in enumerate(piece, 1):
                    _write_text(sheet.cell(row_number, column), value)
            end_row = 8 + len(cutlist.pieces)
            _write_text(sheet.cell(end_row, 1), "EndOfFile")
            for row_number in range(8, end_row + 1):
                sheet.row_dimensions[row_number].height = row_height
                for column, style in enumerate(styles, 1):
                    sheet.cell(row_number, column)._style = copy(style)
            if sheet.max_row > end_row:
                sheet.delete_rows(end_row + 1, sheet.max_row - end_row)
            sheet.print_area = f"A1:K{end_row}"

        input_path = Path(file_path)
        base_name = f"CUTTING - {input_path.stem}"
        suffix = 0
        while True:
            output_path = input_path.with_name(f"{base_name}{f' ({suffix})' if suffix else ''}.xlsx")
            try:
                output = output_path.open("xb")
                break
            except FileExistsError:
                suffix += 1
        try:
            with output:
                workbook.save(output)
        except Exception:
            output_path.unlink()
            raise
        return output_path
    finally:
        workbook.close()
