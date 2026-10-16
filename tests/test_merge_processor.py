from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from merge_processor import group_cutlists, merge_cutlists, move_items, read_merge_file


ROOT = Path(__file__).resolve().parents[1]


def make_export(path, *, job="Job A", board="White", thickness=16, edging="1mm PVC", rows=2):
    workbook = load_workbook(ROOT / "template.xlsx")
    for sheet in list(workbook)[1:]:
        workbook.remove(sheet)
    sheet = workbook.active
    sheet["C2"], sheet["C3"], sheet["J2"] = "Customer", job, "012 345 6789"
    sheet["C4"], sheet["E4"], sheet["J4"] = "Plain Boards", board, thickness
    sheet["C5"], sheet["E5"] = edging, "White"
    for row in range(8, rows + 8):
        for column, value in enumerate([row - 8, 1200 + row, 550, 2, 1, 2, 4, 1,
                                        edging, "White", f"{job}: swing notes"], 1):
            sheet.cell(row, column, value)
    sheet.cell(8 + rows, 1, "EndOfFile")
    workbook.save(path)
    workbook.close()


class MergeTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.first = self.folder / "First file.xlsx"
        self.second = self.folder / "Second file.xlsx"
        self.output = self.folder / "Merged.xlsx"
        make_export(self.first)
        make_export(self.second, job="Job B", rows=3)

    def edit(self, path, callback):
        workbook = load_workbook(path)
        try:
            callback(workbook)
            workbook.save(path)
        finally:
            workbook.close()

    def groups(self, paths=None):
        return group_cutlists([read_merge_file(p) for p in paths or [self.first, self.second]])

    def merge(self, groups=None, **kwargs):
        path = merge_cutlists(groups if groups is not None else self.groups(), self.output,
                              customer="Merged Customer", job="Merged Job", contact="099", **kwargs)
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        return workbook

    def test_matching_boards_combine_without_consolidating_piece_rows(self):
        groups = self.groups()
        self.assertEqual(len(groups), 1)
        self.assertEqual((groups[0].row_count, groups[0].quantity), (5, 10))
        sheet = self.merge(groups).active
        self.assertEqual(sheet["A13"].value, "EndOfFile")
        self.assertEqual([sheet.cell(r, 1).value for r in range(8, 13)], [0, 1, 0, 1, 2])
        self.assertEqual([sheet.cell(r, 11).value for r in range(8, 13)],
                         ["Job A: swing notes"] * 2 + ["Job B: swing notes"] * 3)
        for row in range(8, 13):
            self.assertEqual([sheet.cell(row, c).value for c in range(5, 11)], [1, 2, 4, 1, "1mm PVC", "White"])

    def test_file_order_controls_piece_order(self):
        sheet = self.merge(self.groups([self.second, self.first])).active
        self.assertEqual(sheet["K8"].value, "Job B: swing notes")
        self.assertEqual(sheet["K11"].value, "Job A: swing notes")

    def test_same_board_with_different_edging_combines_and_preserves_per_piece_edging(self):
        make_export(self.second, job="Job B", edging="2mm PVC")
        groups = self.groups()
        self.assertEqual(len(groups), 1)
        sheet = self.merge(groups).active
        self.assertEqual(sheet["I8"].value, "1mm PVC")
        self.assertEqual(sheet["I10"].value, "2mm PVC")

    def test_missing_piece_edging_uses_own_source_header(self):
        make_export(self.second, edging="2mm PVC")
        def clear(w):
            w.active["I8"], w.active["J8"] = None, None
        self.edit(self.second, clear)
        sheet = self.merge().active
        self.assertEqual((sheet["I10"].value, sheet["J10"].value), ("2mm PVC", "White"))

    def test_different_boards_thicknesses_categories_and_settings_stay_separate(self):
        for address, value in [("E4", "Grey"), ("J4", 18), ("C4", "Grain Boards"), ("G6", "FALSE")]:
            with self.subTest(address=address):
                make_export(self.second)
                self.edit(self.second, lambda w: setattr(w.active[address], "value", value))
                self.assertEqual(len(self.groups()), 2)

    def test_case_and_numeric_thickness_representation_do_not_split_board(self):
        make_export(self.second, board=" WHITE ", thickness="16.0")
        self.assertEqual(len(self.groups()), 1)

    def test_final_order_sets_sequential_cutlist_names(self):
        make_export(self.second, board="Grey")
        workbook = self.merge(list(reversed(self.groups())))
        self.assertEqual(workbook.sheetnames, ["CutList1", "CutList2"])
        self.assertEqual(workbook["CutList1"]["E4"].value, "Grey")
        self.assertEqual(workbook["CutList2"]["E4"].value, "White")

    def test_common_details_flags_and_header_layout(self):
        sheet = self.merge().active
        self.assertEqual([sheet[c].value for c in ["C2", "C3", "J2"]], ["Merged Customer", "Merged Job", "099"])
        self.assertEqual([sheet[c].value for c in ["K5", "A6", "G6", "I6"]], ["File Version 7", "FALSE", "TRUE", 0])
        self.assertIn("C2:I2", str(sheet.merged_cells))
        self.assertEqual(sheet["B8"].number_format, "0;\\-0")

    def test_more_than_nine_boards_and_one_hundred_rows(self):
        make_export(self.first, rows=105)
        files = [read_merge_file(self.first)]
        for i in range(10):
            path = self.folder / f"Board {i}.xlsx"
            make_export(path, board=f"Board {i}")
            files.append(read_merge_file(path))
        workbook = self.merge(group_cutlists(files))
        self.assertEqual(len(workbook.worksheets), 11)
        self.assertEqual(workbook.active["A113"].value, "EndOfFile")
        self.assertEqual(workbook.active["B112"]._style, workbook.active["B8"]._style)
        self.assertIn("$K$113", workbook.active.print_area)

    def test_source_bytes_are_unchanged(self):
        before = [p.read_bytes() for p in [self.first, self.second]]
        self.merge()
        self.assertEqual(before, [p.read_bytes() for p in [self.first, self.second]])

    def test_source_cannot_be_overwritten_even_if_requested(self):
        with self.assertRaisesRegex(ValueError, "source files cannot be overwritten"):
            merge_cutlists(self.groups(), self.first, customer="C", job="J", overwrite=True)

    def test_existing_output_requires_explicit_overwrite(self):
        self.output.write_bytes(b"existing")
        with self.assertRaises(FileExistsError):
            self.merge()
        self.assertEqual(self.output.read_bytes(), b"existing")
        self.assertEqual(set(p.name for p in self.folder.iterdir()), {"First file.xlsx", "Second file.xlsx", "Merged.xlsx"})
        self.merge(overwrite=True)
        self.assertTrue(self.output.read_bytes().startswith(b"PK"))

    def test_failed_save_preserves_existing_output_and_cleans_temporary(self):
        self.output.write_bytes(b"existing")
        with patch("openpyxl.workbook.workbook.Workbook.save", side_effect=OSError("disk full")):
            with self.assertRaisesRegex(OSError, "disk full"):
                self.merge(overwrite=True)
        self.assertEqual(self.output.read_bytes(), b"existing")
        self.assertEqual(len(list(self.folder.iterdir())), 3)

    def test_missing_end_marker_is_repaired(self):
        self.edit(self.first, lambda w: setattr(w.active["A10"], "value", None))
        self.assertEqual(self.merge().active["A13"].value, "EndOfFile")

    def test_invalid_formats_versions_rows_and_formulas_are_rejected(self):
        for address, value, message in [("K5", "File Version 8", "Version 7"), ("B8", 0, "row 8"),
                                        ("B8", "=1+2", "formulas"), ("D8", -1, "row 8"),
                                        ("G8", 1.5, "row 8"), ("L8", "special instruction", "beyond"),
                                        ("B7", "wrong", "column headings"), ("J4", "NaN", "thickness")]:
            with self.subTest(address=address, value=value):
                make_export(self.first)
                self.edit(self.first, lambda w: setattr(w.active[address], "value", value))
                with self.assertRaisesRegex(ValueError, message):
                    read_merge_file(self.first)
        with self.assertRaisesRegex(ValueError, "xlsx export"):
            read_merge_file(self.folder / "legacy.xls")
        make_export(self.first)
        self.edit(self.first, lambda w: setattr(w.active["A1"], "value", "Cupboard Manager 3D Cutting List"))
        with self.assertRaisesRegex(ValueError, "not a Cut Manager"):
            read_merge_file(self.first)

    def test_invalid_output_and_missing_details(self):
        with self.assertRaisesRegex(ValueError, "customer and job"):
            merge_cutlists(self.groups(), self.output, customer="", job="J")
        with self.assertRaisesRegex(ValueError, "xlsx extension"):
            merge_cutlists(self.groups(), self.output.with_suffix(".xls"), customer="C", job="J")


class OrderingTests(unittest.TestCase):
    def test_move_multiple_rows_preserves_relative_order_and_selection(self):
        self.assertEqual(move_items(list("ABCDE"), [1, 2], -1), (list("BCADE"), [0, 1]))
        self.assertEqual(move_items(list("ABCDE"), [1, 2], 1), (list("ADBCE"), [2, 3]))

    def test_boundary_and_disjoint_selection(self):
        self.assertEqual(move_items(list("ABCDE"), [0, 2], -1), (list("ACBDE"), [0, 1]))
        self.assertEqual(move_items(list("ABCDE"), [0, 4], 1), (list("BACDE"), [1, 4]))
        self.assertEqual(move_items(list("ABC"), [], -1), (list("ABC"), []))


if __name__ == "__main__":
    unittest.main()
