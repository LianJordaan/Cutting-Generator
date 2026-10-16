from contextlib import ExitStack, redirect_stdout
from io import StringIO
from pathlib import Path
import os
import runpy
import shutil
import subprocess
import sys
from tempfile import TemporaryDirectory
from types import ModuleType, SimpleNamespace
import unittest
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from config_utils import DEVELOPER_PASSWORD
from cupboard_processor import is_cupboard_cutlist, process_cupboard_cutlist


ROOT = Path(__file__).resolve().parents[1]
HEADERS = [
    "Unit ID", "L", "W", "Qty", "Board Category", "Board Name",
    "Edge L", "Edge W", "Hole L", "Hole W", "Edging Category", "Edging Name", "Remarks",
]
BOARD_ROW = [0, 2399, 550, 2, "Plain Boards", "Peen White", 1, 2, 4, 1,
             "1mm PVC", "White", "1xL Swing. 1xR Swing."]
HARDBOARD_ROW = [9999, 2300, 971, 1, "Plain Hardboard", "White HB 9x6",
                0, 0, 0, 0, "1mm PVC", "White", "Left Side"]


def make_source(path, *, board_rows=2, sheet_count=2, start_row=6):
    workbook = Workbook()
    info = workbook.active
    info.title = "Cutlist Info"
    info["A2"] = "Cupboard Manager 3D Cutting List"
    info["B11"] = 11
    info["D20"] = start_row
    info["D21"] = 14
    info["D22"] = 16
    info["D23"] = 18
    info["D24"] = 20
    for index in range(sheet_count):
        sheet = workbook.create_sheet(f"CutList{index + 1}")
        piece = BOARD_ROW if index == 0 else HARDBOARD_ROW
        sheet["A1"] = "Cupboard Manager 3D Cutting List - Plan 1"
        sheet["C2"], sheet["E2"], sheet["K2"] = piece[4], piece[5], "16" if index == 0 else "3"
        sheet["C3"] = "1mm PVC" if index == 0 else None
        sheet["E3"] = "White"
        for column, heading in enumerate(HEADERS, 1):
            sheet.cell(start_row - 1, column, heading)
        count = board_rows if index == 0 else 1
        for row in range(start_row, start_row + count):
            for column, value in enumerate(piece, 1):
                sheet.cell(row, column, value)
        sheet.cell(start_row + count, 1, "EndOfFile")
    workbook.save(path)
    workbook.close()


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.source = self.folder / "299 - Example Customer_Job-Kamerkaste_Plan-Plan 1.xlsx"
        make_source(self.source)

    def edit(self, callback):
        workbook = load_workbook(self.source)
        try:
            callback(workbook)
            workbook.save(self.source)
        finally:
            workbook.close()

    def convert(self):
        return process_cupboard_cutlist(self.source, ROOT / "template.xlsx")

    def read_output(self):
        workbook = load_workbook(self.convert())
        self.addCleanup(workbook.close)
        return workbook

    def assert_invalid(self, message):
        before = set(self.folder.iterdir())
        with self.assertRaisesRegex(ValueError, message):
            self.convert()
        self.assertEqual(before, set(self.folder.iterdir()))

    def test_detects_content_with_arbitrary_filename_and_future_version(self):
        self.edit(lambda w: setattr(w["Cutlist Info"]["B11"], "value", 12))
        renamed = self.folder / "anything.xlsx"
        shutil.copyfile(self.source, renamed)
        self.assertTrue(is_cupboard_cutlist(renamed))

    def test_does_not_detect_cut_manager_or_erik(self):
        self.assertFalse(is_cupboard_cutlist(ROOT / "template.xlsx"))
        self.assertFalse(is_cupboard_cutlist(self.folder / "erik.xls"))

    def test_maps_pieces_metadata_and_hardboard_without_renaming(self):
        workbook = self.read_output()
        self.assertEqual(workbook.sheetnames, ["CutList1", "CutList2"])
        sheet = workbook["CutList1"]
        self.assertEqual([sheet[c].value for c in ["C2", "C3", "C4", "E4", "J4", "C5", "E5"]],
                         ["Example Customer", "Kamerkaste", "Plain Boards", "Peen White", 16, "1mm PVC", "White"])
        expected = [0, 2399, 550, 2, 1, 2, 4, 1, "1mm PVC", "White", None]
        self.assertEqual([c.value for c in sheet[8]], expected)
        self.assertEqual([c.value for c in sheet[9]], expected)
        self.assertEqual(sheet["A10"].value, "EndOfFile")
        hardboard = workbook["CutList2"]
        self.assertEqual([hardboard[c].value for c in ["C4", "E4", "J4", "C5", "E5"]],
                         ["Plain Hardboard", "White HB 9x6", 3, None, "White"])
        self.assertEqual([c.value for c in hardboard[8]],
                         [9999, 2300, 971, 1, 0, 0, 0, 0, "1mm PVC", "White", None])

    def test_preserves_template_structure_and_control_flags(self):
        workbook = self.read_output()
        template = load_workbook(ROOT / "template.xlsx")
        self.addCleanup(template.close)
        for sheet in workbook:
            self.assertEqual(str(sheet.merged_cells), str(template.active.merged_cells))
            for address in ["A1", "K2", "K5", "A6", "C6", "E6", "G6", "I6", "K6"]:
                self.assertEqual(sheet[address].value, template.active[address].value)
            self.assertEqual([c.value for c in sheet[7]], [c.value for c in template.active[7]])

    def test_prefers_info_customer_and_contact(self):
        def change(w):
            w["Cutlist Info"]["D13"] = "Actual Customer"
            w["Cutlist Info"]["D14"] = "012 345 6789"
        self.edit(change)
        workbook = self.read_output()
        for sheet in workbook:
            self.assertEqual(sheet["C2"].value, "Actual Customer")
            self.assertEqual(sheet["J2"].value, "012 345 6789")

    def test_extends_rows_and_adds_sheets_without_losing_data(self):
        make_source(self.source, board_rows=105, sheet_count=11)
        workbook = self.read_output()
        self.assertEqual(len(workbook.worksheets), 11)
        sheet = workbook.active
        self.assertEqual(sheet["B112"].value, 2399)
        self.assertEqual(sheet["A113"].value, "EndOfFile")
        self.assertEqual(sheet["B112"]._style, sheet["B8"]._style)
        self.assertEqual(sheet["K112"]._style, sheet["K8"]._style)
        self.assertIn("$A$1:$K$113", sheet.print_area)
        self.assertEqual(workbook["CutList11"]["A9"].value, "EndOfFile")

    def test_uses_declared_start_row(self):
        make_source(self.source, start_row=8)
        self.assertEqual(self.read_output().active["B8"].value, 2399)

    def test_missing_end_marker_and_empty_rows(self):
        def change(w):
            s = w["CutList1"]
            s["A8"] = None
            s.insert_rows(7)
        self.edit(change)
        sheet = self.read_output().active
        self.assertEqual(sheet["A10"].value, "EndOfFile")
        self.assertEqual(sheet.max_row, 10)

    def test_numeric_strings_and_fractional_dimensions_are_preserved(self):
        def change(w):
            s = w["CutList1"]
            s["B6"], s["C6"], s["D6"] = "2399,5", 550.25, "2"
        self.edit(change)
        sheet = self.read_output().active
        self.assertEqual([sheet[c].value for c in ["B8", "C8", "D8"]], [2399.5, 550.25, 2])

    def test_source_names_beginning_with_equals_remain_literal(self):
        self.edit(lambda w: setattr(w["Cutlist Info"]["D13"], "value", "=Customer"))
        cell = self.read_output().active["C2"]
        self.assertEqual((cell.value, cell.data_type), ("=Customer", "s"))

    def test_output_collisions_preserve_existing_files_and_input(self):
        original = self.source.read_bytes()
        first = self.convert()
        first_bytes = first.read_bytes()
        second = self.convert()
        self.assertEqual(first.name, f"CUTTING - {self.source.stem}.xlsx")
        self.assertEqual(second.name, f"CUTTING - {self.source.stem} (1).xlsx")
        self.assertEqual(first.read_bytes(), first_bytes)
        self.assertEqual(self.source.read_bytes(), original)

    def test_removes_partial_output_if_save_fails(self):
        with patch("openpyxl.workbook.workbook.Workbook.save", side_effect=OSError("disk full")):
            with self.assertRaisesRegex(OSError, "disk full"):
                self.convert()
        self.assertEqual(list(self.folder.iterdir()), [self.source])

    def test_unsupported_version_is_rejected_before_output(self):
        self.edit(lambda w: setattr(w["Cutlist Info"]["B11"], "value", 12))
        self.assert_invalid("Cutlist Info, row 11: unsupported file version")

    def test_missing_info_sheet_is_still_detected_then_rejected(self):
        self.edit(lambda w: w.remove(w["Cutlist Info"]))
        self.assertTrue(is_cupboard_cutlist(self.source))
        self.assert_invalid("missing its Cutlist Info")

    def test_unusable_filename_is_rejected(self):
        self.source = self.source.rename(self.folder / "arbitrary.xlsx")
        self.assert_invalid("Cannot determine customer/job")

    def test_invalid_rows_are_reported_without_output(self):
        for address, value in [("B6", "bad"), ("C6", 0), ("D6", 1.5), ("G6", 3),
                               ("I6", -1), ("A6", None), ("B6", "=1+2"), ("B6", "NaN")]:
            with self.subTest(address=address, value=value):
                make_source(self.source)
                self.edit(lambda w: setattr(w["CutList1"][address], "value", value))
                self.assert_invalid("CutList1, row 6: invalid")

    def test_unsupported_operations_anywhere_are_rejected(self):
        for column in [14, 15, 16, 17, 18, 19, 20, 21]:
            with self.subTest(column=column):
                make_source(self.source)
                self.edit(lambda w: w["CutList2"].cell(6, column, "instruction"))
                self.assert_invalid("CutList2, row 6:.*unsupported")

    def test_invalid_header_or_board_mismatch_is_rejected(self):
        self.edit(lambda w: setattr(w["CutList1"]["F6"], "value", "Other Board"))
        self.assert_invalid("CutList1, row 6:.*differs")
        make_source(self.source)
        self.edit(lambda w: setattr(w["CutList1"]["B5"], "value", "Wrong"))
        self.assert_invalid("CutList1, row 5: expected")

    def test_data_after_end_marker_is_rejected(self):
        self.edit(lambda w: w["CutList2"].cell(8, 1, "unexpected"))
        self.assert_invalid("CutList2, row 8: unexpected data after EndOfFile")

    def test_empty_cutlists_and_missing_board_heading_are_rejected(self):
        make_source(self.source, board_rows=0)
        self.assert_invalid("CutList1, row 6: cutlist contains no pieces")
        make_source(self.source)
        self.edit(lambda w: setattr(w["CutList2"]["E2"], "value", None))
        self.assert_invalid("CutList2, row 2: board category and board name are required")


def fake_module(name, **values):
    module = ModuleType(name)
    module.__dict__.update(values)
    return module


class RoutingTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / "123 - Customer_Job-Job_Plan-Plan 1.xlsx"
        make_source(self.source)
        self.config = dict(ip="localhost", username="test", password="test", agree_terms=True,
                           developer_password=DEVELOPER_PASSWORD)
        self.setup_info = Mock(side_effect=lambda: self.config)
        self.license = Mock(return_value=SimpleNamespace(allowed=True, checked_remotely=False))
        self.convert = Mock(wraps=process_cupboard_cutlist)
        self.regular = Mock(return_value=[])
        self.erik = Mock()
        self.erik_detector = Mock(return_value=False)
        self.cutouts = Mock(return_value=[])
        self.merge = Mock(return_value=None)
        self.modules = {
            "setup_gui": fake_module("setup_gui", setup=Mock(), get_setup_info=self.setup_info,
                                     show_license_issue_window=Mock()),
            "excel_processor": fake_module("excel_processor", process_excel=self.regular,
                                           process_erik_cutlist=self.erik, get_job_name=Mock(return_value="Job"),
                                           find_cutouts=self.cutouts, find_crosscuts=Mock(return_value=[]),
                                           containsSpecialCrosscuts=Mock(return_value=False)),
            "helpers": fake_module("helpers", is_erik_cutlist=self.erik_detector,
                                   refresh_desktop=Mock(), find_quote_num=Mock(return_value="123")),
            "shape_gen": fake_module("shape_gen"),
            "device_auth": fake_module("device_auth", ensure_device_is_authorized=self.license,
                                       get_machine_guid=Mock()),
        }

    def run_app(self, path=None, paths=None):
        output = StringIO()
        with ExitStack() as stack:
            stack.enter_context(patch.dict(sys.modules, self.modules))
            stack.enter_context(patch.object(sys, "argv", [str(ROOT / "app.py"), *(paths or [str(path or self.source)])]))
            stack.enter_context(patch("os.system"))
            stack.enter_context(patch("builtins.input", return_value=""))
            stack.enter_context(patch("cupboard_processor.process_cupboard_cutlist", self.convert))
            stack.enter_context(patch("merge_gui.show_merge_window", self.merge))
            stack.enter_context(redirect_stdout(output))
            try:
                runpy.run_path(str(ROOT / "app.py"), run_name="__main__")
            except SystemExit as exc:
                return exc.code, output.getvalue()
        return None, output.getvalue()

    def test_missing_and_wrong_password_cannot_convert_or_fall_through(self):
        for password in [None, "", "wrong"]:
            with self.subTest(password=password):
                if password is None:
                    self.config.pop("developer_password", None)
                else:
                    self.config["developer_password"] = password
                code, output = self.run_app()
                self.assertEqual(code, 1)
                self.assertIn("requires the Developer Password in Setup", output)
        self.convert.assert_not_called()
        self.regular.assert_not_called()
        self.erik.assert_not_called()
        self.assertEqual(list(self.source.parent.iterdir()), [self.source])

    def test_correct_password_creates_workbook_and_finishes(self):
        code, output = self.run_app()
        self.assertEqual(code, 0)
        self.assertIn("Saved Cut Manager workbook:", output)
        self.license.assert_called_once()
        self.convert.assert_called_once()
        self.regular.assert_not_called()
        self.cutouts.assert_not_called()
        self.assertTrue(self.source.with_name(f"CUTTING - {self.source.name}").exists())

    def test_invalid_licence_stops_before_setup_and_conversion(self):
        self.license.return_value = SimpleNamespace(allowed=False, message="Licence required", machine_guid="test")
        code, output = self.run_app()
        self.assertEqual(code, 0)
        self.assertIn("Licence required", output)
        self.setup_info.assert_not_called()
        self.convert.assert_not_called()

    def test_ordinary_cut_manager_file_uses_existing_processor(self):
        self.config.pop("developer_password")
        code, _ = self.run_app(ROOT / "test.xlsx")
        self.assertIsNone(code)
        self.regular.assert_called_once()
        self.convert.assert_not_called()

    def test_erik_route_unchanged(self):
        self.erik_detector.return_value = True
        code, _ = self.run_app(self.source.with_suffix(".xls"))
        self.assertEqual(code, 0)
        self.erik.assert_called_once()
        self.regular.assert_not_called()
        self.convert.assert_not_called()

    def test_validation_failure_does_not_fall_through(self):
        self.convert.side_effect = ValueError("unsupported file version")
        _, output = self.run_app()
        self.assertIn("unsupported file version", output)
        self.assertIn("Processing aborted", output)
        self.regular.assert_not_called()

    def test_multiple_files_open_merge_window_and_do_not_process_single_file(self):
        paths = [str(self.source), str(self.source.with_name("second.xlsx"))]
        code, _ = self.run_app(paths=paths)
        self.assertEqual(code, 0)
        self.merge.assert_called_once_with(paths, config=self.config)
        self.regular.assert_not_called()
        self.convert.assert_not_called()
        self.erik.assert_not_called()

    def test_merge_requires_developer_password(self):
        for password in [None, "", "wrong"]:
            with self.subTest(password=password):
                self.config["developer_password"] = password
                code, output = self.run_app(paths=["first.xlsx", "second.xlsx"])
                self.assertEqual(code, 1)
                self.assertIn("requires the Developer Password", output)
        self.merge.assert_not_called()

    def test_transfer_record_resumes_only_after_developer_check(self):
        path = Path(self.temp.name) / "Merged.transfer.json"
        with patch("transfer_gui.resume_transfer") as resume:
            code, _ = self.run_app(path=path)
            self.assertEqual(code, 0)
            resume.assert_called_once_with(str(path), self.config)
            resume.reset_mock()
            self.config["developer_password"] = "wrong"
            code, _ = self.run_app(path=path)
            self.assertEqual(code, 1)
            resume.assert_not_called()
        self.regular.assert_not_called()
        self.regular.assert_not_called()

    def test_unlicensed_machine_cannot_open_merge_window(self):
        self.license.return_value = SimpleNamespace(allowed=False, message="Licence required", machine_guid="test")
        self.run_app(paths=["first.xlsx", "second.xlsx"])
        self.merge.assert_not_called()

    def test_completed_merge_reports_saved_path(self):
        self.merge.return_value = Path(self.temp.name) / "merged.xlsx"
        code, output = self.run_app(paths=["first.xlsx", "second.xlsx"])
        self.assertEqual(code, 0)
        self.assertIn("Saved merged Cut Manager workbook:", output)

    @unittest.skipUnless(sys.platform == "win32", "Machine GUID is Windows-only")
    def test_startup_handles_redirected_windows_text_encoding(self):
        result = subprocess.run(
            [sys.executable, str(ROOT / "app.py"), "--machine-guid"],
            capture_output=True, encoding="cp1252", timeout=30,
            env={**os.environ, "PYTHONIOENCODING": "cp1252", "PYTHONUTF8": "0"},
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Libraries loaded.", result.stdout)


if __name__ == "__main__":
    unittest.main()
