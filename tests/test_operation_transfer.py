"""Integration tests use newly created, disposable Firebird databases ONLY."""

from copy import deepcopy
from datetime import datetime, timezone, timedelta
import json
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

import fdb
from openpyxl import load_workbook

from config_utils import DEVELOPER_PASSWORD
from merge_processor import group_cutlists, merge_cutlists, read_merge_file
import operation_transfer as transfer
from test_merge_processor import make_export


LIBRARY = Path(r"C:\Program Files\Firebird\Firebird_3_0\fbclient.dll")
SCHEMA = [
    "CREATE TABLE DB_INFO (VERSION INTEGER)",
    "CREATE TABLE CUSTOMER (CUSTOMER_ID INTEGER, NAME VARCHAR(100))",
    "CREATE TABLE QUOTE (QUOTE_NR INTEGER, JOB_NAME VARCHAR(100), CUSTOMER_ID INTEGER, SIZEMODE SMALLINT)",
    "CREATE TABLE CUTLIST (QUOTE_NR INTEGER,CUTLIST_ID INTEGER,STATUS SMALLINT,OPTIMIZE_MODE SMALLINT,BUILD_TYPE VARCHAR(1))",
    "CREATE TABLE BOARD_CATEGORY (BCAT_ID INTEGER,BCAT_NAME VARCHAR(100))",
    "CREATE TABLE BOARD_TYPES (BOARD_ID INTEGER,BCAT_ID INTEGER,BOARD_NAME VARCHAR(100),BOARD_THICKNESS INTEGER)",
    "CREATE TABLE EDGE_CATEGORY (ECAT_ID INTEGER,ECAT_NAME VARCHAR(100))",
    "CREATE TABLE EDGE_TYPES (EDGE_ID INTEGER,ECAT_ID INTEGER,EDGE_NAME VARCHAR(100))",
    """CREATE TABLE CUT_LIST_DETAIL (QUOTE_NR INTEGER,CUTLIST_ID INTEGER,ITEM_ID INTEGER,UNIT_ID INTEGER,
       LENGTE INTEGER,WYDTE INTEGER,QTY INTEGER,EDGEL INTEGER,EDGEW INTEGER,HOLESL INTEGER,HOLESW INTEGER,
       REMARKS VARCHAR(100),BOARD_ID INTEGER,EDGE_ID INTEGER)""",
    "CREATE TABLE CROSSCUTS (QUOTE_NR INTEGER,CUTLIST_ID INTEGER,ITEM_ID INTEGER,CC_ID INTEGER,LENGTE INTEGER,SPLIT INTEGER,HOLES_L INTEGER,HOLES_S INTEGER,LSWING VARCHAR(1))",
    "CREATE TABLE CUTOUTS (QUOTE_NR INTEGER,CUTLIST_ID INTEGER,ITEM_ID INTEGER,CUTOUT1 VARCHAR(100),CUTOUT2 VARCHAR(100))",
    "CREATE TABLE MITRES (QUOTE_NR INTEGER,CUTLIST_ID INTEGER,ITEM_ID INTEGER,MITRE_L VARCHAR(100),MITRE_W VARCHAR(100))",
    "CREATE TABLE GROOVES (QUOTE_NR INTEGER,CUTLIST_ID INTEGER,ITEM_ID INTEGER,GROOVE_L VARCHAR(100),GROOVE_W VARCHAR(100))",
]


@unittest.skipUnless(sys.platform == "win32" and LIBRARY.is_file(), "Requires local Firebird 3 embedded engine")
class TransferIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory(prefix="cutting-transfer-test-")
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name).resolve()
        self.database = self.folder / "disposable.fdb"
        self.config = {"ip": "", "filepath": str(self.database), "username": "SYSDBA",
                       "password": "fixture-only", "charset": "UTF8", "fb_library": str(LIBRARY),
                       "developer_password": DEVELOPER_PASSWORD}
        self.key_patch = patch("operation_transfer._key", return_value=b"disposable-test-signing-key")
        self.key_patch.start()
        self.addCleanup(self.key_patch.stop)
        connection = fdb.create_database(dsn=str(self.database), user="SYSDBA", password="fixture-only", fb_library_name=str(LIBRARY))
        try:
            cursor = connection.cursor()
            for statement in SCHEMA:
                cursor.execute(statement)
            connection.commit()
        finally:
            connection.close()
        self.sql("INSERT INTO DB_INFO VALUES (24)")
        self.sql("INSERT INTO CUSTOMER VALUES (1,'Customer')")
        self.sql("INSERT INTO BOARD_CATEGORY VALUES (1,'Plain Boards')")
        self.sql("INSERT INTO BOARD_TYPES VALUES (1,1,'White',16)")
        self.sql("INSERT INTO EDGE_CATEGORY VALUES (1,'1mm PVC')")
        self.sql("INSERT INTO EDGE_TYPES VALUES (1,1,'White')")
        self.paths = []
        for number, job in [(100, "Job A"), (101, "Job B")]:
            path = self.folder / f"{number} - Customer_Job-{job}.xlsx"
            make_export(path, job=job)
            self.paths.append(path)
            self.sql("INSERT INTO QUOTE VALUES (?,?,1,1)", (number, job))
            self.sql("INSERT INTO CUTLIST VALUES (?,1,0,0,'3')", (number,))
            workbook = load_workbook(path)
            try:
                for row in [8, 9]:
                    values = [workbook.active.cell(row, c).value for c in range(1, 12)]
                    self.insert_piece(number, 1, row - 7, values)
            finally:
                workbook.close()
        self.sql("INSERT INTO CROSSCUTS VALUES (100,1,1,0,185,0,2,0,'0')")
        self.sql("INSERT INTO CUTOUTS VALUES (100,1,2,'3205000500','')")
        self.sql("INSERT INTO GROOVES VALUES (101,1,1,'1','0')")
        self.sql("INSERT INTO MITRES VALUES (101,1,2,'11','00')")
        self.files = [read_merge_file(path) for path in self.paths]
        self.groups = group_cutlists(self.files)

    def connect(self):
        # Absolute path under TemporaryDirectory; no host and no real app config.
        self.assertTrue(self.database.is_relative_to(self.folder))
        return fdb.connect(database=str(self.database), user="SYSDBA", password="fixture-only", fb_library_name=str(LIBRARY))

    def sql(self, statement, parameters=()):
        connection = self.connect()
        try:
            cursor = connection.cursor()
            cursor.execute(statement, parameters)
            result = cursor.fetchall() if statement.startswith("SELECT") else None
            connection.commit()
            return result
        finally:
            connection.close()

    def insert_piece(self, quote, cutlist, item, values):
        self.sql("INSERT INTO CUT_LIST_DETAIL VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                 (quote, cutlist, item, *values[:8], values[10], 1, 1))

    def stage(self):
        self.inspection = transfer.inspect_sources(self.files, self.config)
        self.assertFalse(self.inspection["errors"], self.inspection["errors"])
        payload, unit_ids = transfer.prepare_manifest(self.groups, self.inspection, "Customer", "Merged Job")
        self.workbook_path = self.folder / "Merged.xlsx"
        merge_cutlists(self.groups, self.workbook_path, customer="Customer", job="Merged Job", unit_ids=unit_ids)
        self.manifest_path = transfer.save_manifest(payload, self.workbook_path)
        self.sql("INSERT INTO QUOTE VALUES (102,'Merged Job',1,1)")
        self.sql("INSERT INTO CUTLIST VALUES (102,1,0,0,'3')")
        workbook = load_workbook(self.workbook_path)
        try:
            for row in range(8, 12):
                self.insert_piece(102, 1, 500 + row, [workbook.active.cell(row, c).value for c in range(1, 12)])
        finally:
            workbook.close()
        self.plan = transfer.check_target(self.manifest_path, 102, self.config)
        return self.plan

    def backup(self):
        return transfer.create_verified_backup(self.manifest_path, self.plan, self.config, self.folder / "backups")

    def apply(self, receipt, *, reviewed=False):
        return transfer.apply_transfer(self.manifest_path, self.plan, receipt, self.config,
                                       confirmed_quote=102, cutting_manager_closed=True,
                                       confirmed_review_hash=transfer.digest(self.plan["review_differences"]) if reviewed else None)

    def operation_rows(self, quote=102):
        return {table: self.sql(f"SELECT * FROM {table} WHERE QUOTE_NR=?", (quote,)) for table in transfer.OPERATIONS}

    def test_readonly_discovery_finds_all_four_operation_types(self):
        inspection = transfer.inspect_sources(self.files, self.config)
        self.assertFalse(inspection["errors"])
        self.assertEqual(transfer.operation_counts(inspection), {"CROSSCUTS": 1, "CUTOUTS": 1, "MITRES": 1, "GROOVES": 1})

    def test_unique_ids_preserve_original_ids_in_signed_manifest(self):
        self.stage()
        manifest = transfer.read_manifest(self.manifest_path, self.config)
        self.assertEqual([r["uid"] for r in manifest["rows"]], [1, 2, 3, 4])
        self.assertEqual([r["original_unit_id"] for r in manifest["rows"]], [0, 1, 0, 1])
        self.assertFalse(any(self.operation_rows().values()))

    def rewrite_workbook(self, change):
        workbook = load_workbook(self.workbook_path)
        try:
            change(workbook)
            workbook.save(self.workbook_path)
        finally:
            workbook.close()

    def rewrite_imported_edges(self, workbook):
        for sheet in workbook:
            sheet["E5"] = "EID_1"
            for row in range(8, 12):
                sheet.cell(row, 10).value = "EID_1"

    def test_importer_edge_ids_and_resaving_allow_verified_transfer(self):
        self.stage()
        self.rewrite_workbook(self.rewrite_imported_edges)
        before = self.workbook_path.read_bytes()
        self.assertEqual(transfer.check_target(self.manifest_path, 102, self.config), self.plan)
        self.assertEqual(self.workbook_path.read_bytes(), before, "Verification must not rewrite the user's workbook")
        self.assertEqual(self.apply(self.backup()), "committed")
        self.assertEqual(sum(len(rows) for rows in self.operation_rows().values()), 4)

    def test_existing_v530_records_accept_importer_edge_id_rewrites(self):
        self.stage()
        payload = json.loads(self.manifest_path.read_text())["payload"]
        payload.pop("workbook_content_hash")
        transfer._write_signed(self.manifest_path, payload)
        self.rewrite_workbook(self.rewrite_imported_edges)
        self.assertEqual(transfer.read_manifest(self.manifest_path, self.config), payload)
        self.assertEqual(len(transfer.check_target(self.manifest_path, 102, self.config)["inserts"]), 4)

    def test_document_metadata_only_resave_requires_no_database_access(self):
        self.stage()
        self.rewrite_workbook(lambda w: setattr(w.properties, "lastModifiedBy", "Cut Manager import"))
        with patch("operation_transfer._connect", side_effect=AssertionError("Unexpected database access")):
            manifest = transfer.read_manifest(self.manifest_path, self.config)
        self.assertNotEqual(transfer.file_hash(self.workbook_path), manifest["workbook_hash"])

    def test_real_piece_heading_flag_and_structure_changes_still_fail(self):
        self.stage()
        original = self.workbook_path.read_bytes()
        edits = [("B8", 999), ("A8", 999), ("D8", 3), ("G8", 99),
                 ("K8", "Changed instructions"), ("J8", "Other edging"),
                 ("E4", "Other board"), ("C2", "Other customer"),
                 ("C3", "Other job"), ("A6", "Changed flag"),
                 ("A12", None)]
        for address, value in edits:
            with self.subTest(address=address):
                self.workbook_path.write_bytes(original)
                self.rewrite_workbook(lambda w: setattr(w.active[address], "value", value))
                with self.assertRaises(transfer.TransferError):
                    transfer.read_manifest(self.manifest_path, self.config)
        self.workbook_path.write_bytes(original)
        self.rewrite_workbook(lambda w: w.create_sheet("Unexpected"))
        with self.assertRaises(transfer.TransferError):
            transfer.read_manifest(self.manifest_path, self.config)
        self.assertFalse(any(self.operation_rows().values()))

    def test_different_edge_id_is_rejected_even_with_same_name(self):
        self.stage()
        self.sql("INSERT INTO EDGE_TYPES VALUES (2,1,'White')")
        self.rewrite_workbook(lambda w: setattr(w.active["J8"], "value", "EID_2"))
        with self.assertRaisesRegex(transfer.TransferError, "edging ID differs"):
            transfer.read_manifest(self.manifest_path, self.config)
        self.assertFalse(any(self.operation_rows().values()))

    def test_importer_unused_edge_changes_on_unedged_pieces_are_harmless(self):
        for path in self.paths:
            workbook = load_workbook(path)
            try:
                for row in [8, 9]:
                    workbook.active.cell(row, 5).value = 0
                    workbook.active.cell(row, 6).value = 0
                workbook.save(path)
            finally:
                workbook.close()
        self.sql("UPDATE CUT_LIST_DETAIL SET EDGEL=0,EDGEW=0")
        self.sql("INSERT INTO EDGE_TYPES VALUES (2,1,'Unused import default')")
        self.files = [read_merge_file(path) for path in self.paths]
        self.groups = group_cutlists(self.files)
        self.stage()
        def unused_edges(workbook):
            for row in range(8, 12):
                workbook.active.cell(row, 10).value = "EID_2"
        self.rewrite_workbook(unused_edges)
        self.assertEqual(transfer.check_target(self.manifest_path, 102, self.config), self.plan)

    def test_missing_workbook_error_identifies_its_expected_location(self):
        self.stage()
        self.workbook_path.unlink()
        with self.assertRaisesRegex(transfer.TransferError, "not found:.*Merged.xlsx"):
            transfer.read_manifest(self.manifest_path, self.config)

    def test_verified_backup_and_atomic_transfer_preserve_all_existing_data(self):
        self.stage()
        sources_before = [self.operation_rows(100), self.operation_rows(101)]
        details_before = self.sql("SELECT * FROM CUT_LIST_DETAIL ORDER BY QUOTE_NR,ITEM_ID")
        receipt = self.backup()
        self.assertEqual(self.apply(receipt), "committed")
        self.assertEqual(sources_before, [self.operation_rows(100), self.operation_rows(101)])
        self.assertEqual(details_before, self.sql("SELECT * FROM CUT_LIST_DETAIL ORDER BY QUOTE_NR,ITEM_ID"))
        operations = self.operation_rows()
        self.assertEqual(operations["CROSSCUTS"][0][3:], (0, 185, 0, 2, 0, "0"))
        self.assertEqual(sum(len(rows) for rows in operations.values()), 4)
        self.assertTrue(transfer.check_target(self.manifest_path, 102, self.config)["already_applied"])
        self.assertEqual(self.apply(receipt), "already_applied")
        self.assertEqual(sum(len(rows) for rows in self.operation_rows().values()), 4)

    def test_failure_after_first_insert_rolls_back_every_addition(self):
        self.stage()
        receipt = self.backup()
        real_connect = transfer._connect

        class FailingCursor:
            def __init__(self, cursor):
                self.cursor = cursor
                self.inserts = 0
            def execute(self, sql, values=()):
                if sql.startswith("INSERT"):
                    self.inserts += 1
                    if self.inserts == 2:
                        raise OSError("simulated insertion failure")
                return self.cursor.execute(sql, values)
            def __getattr__(self, name):
                return getattr(self.cursor, name)

        class Connection:
            def __init__(self, underlying):
                self.underlying = underlying
            def cursor(self):
                return FailingCursor(self.underlying.cursor())
            def __getattr__(self, name):
                return getattr(self.underlying, name)

        with patch("operation_transfer._connect", side_effect=lambda config, write=False: Connection(real_connect(config, write=write))):
            with self.assertRaisesRegex(OSError, "insertion failure"):
                self.apply(receipt)
        self.assertFalse(any(self.operation_rows().values()))

    def test_changed_target_after_review_is_rejected_without_writes(self):
        self.stage()
        receipt = self.backup()
        self.sql("UPDATE CUT_LIST_DETAIL SET LENGTE=999 WHERE QUOTE_NR=102 AND UNIT_ID=1")
        with self.assertRaisesRegex(transfer.TransferError, "does not match"):
            self.apply(receipt)
        self.assertFalse(any(self.operation_rows().values()))

    def test_partial_or_conflicting_target_operations_are_never_overwritten(self):
        self.stage()
        self.sql("INSERT INTO CROSSCUTS VALUES (102,1,508,0,123,0,0,0,'0')")
        with self.assertRaisesRegex(transfer.TransferError, "different or partial"):
            transfer.check_target(self.manifest_path, 102, self.config)
        self.assertEqual(self.operation_rows()["CROSSCUTS"][0][4], 123)

    def test_missing_or_changed_backup_blocks_writes(self):
        self.stage()
        receipt = self.backup()
        backup = Path(json.loads(Path(receipt).read_text())["payload"]["path"])
        backup.write_bytes(b"changed")
        with self.assertRaisesRegex(transfer.TransferError, "expired or changed"):
            self.apply(receipt)
        self.assertFalse(any(self.operation_rows().values()))

    def test_source_existing_or_optimized_target_cannot_be_modified(self):
        self.stage()
        with self.assertRaisesRegex(transfer.TransferError, "existing/source"):
            transfer.check_target(self.manifest_path, 100, self.config)
        self.sql("UPDATE CUTLIST SET STATUS=1 WHERE QUOTE_NR=102")
        with self.assertRaisesRegex(transfer.TransferError, "not optimized"):
            transfer.check_target(self.manifest_path, 102, self.config)

    def test_source_changed_after_merge_stops_transfer(self):
        self.stage()
        self.sql("UPDATE CROSSCUTS SET LENGTE=188 WHERE QUOTE_NR=100")
        with self.assertRaisesRegex(transfer.TransferError, "Source quote 100 changed"):
            transfer.check_target(self.manifest_path, 102, self.config)

    def test_edited_manifest_workbook_or_database_identity_are_rejected(self):
        self.stage()
        bad_config = {**self.config, "filepath": str(self.folder / "different.fdb")}
        with self.assertRaisesRegex(transfer.TransferError, "configured database"):
            transfer.read_manifest(self.manifest_path, bad_config)
        envelope = json.loads(self.manifest_path.read_text())
        envelope["payload"]["job"] = "Other Job"
        self.manifest_path.write_text(json.dumps(envelope))
        with self.assertRaisesRegex(transfer.TransferError, "was changed"):
            transfer.read_manifest(self.manifest_path, self.config)

    def test_ambiguous_source_is_blocked_instead_of_guessing(self):
        path = self.paths[0]
        workbook = load_workbook(path)
        try:
            for column in range(1, 12):
                workbook.active.cell(9, column).value = workbook.active.cell(8, column).value
            workbook.save(path)
        finally:
            workbook.close()
        self.sql("UPDATE CUT_LIST_DETAIL SET UNIT_ID=0,LENGTE=1208 WHERE QUOTE_NR=100 AND ITEM_ID=2")
        inspection = transfer.inspect_sources([read_merge_file(path)], self.config)
        self.assertIn("ambiguous", inspection["errors"][str(path)])

    def test_missing_source_quote_is_unknown_not_no_operations(self):
        inspection = transfer.inspect_sources(self.files, self.config, {str(self.paths[0]): 999})
        self.assertIn(str(self.paths[0]), inspection["errors"])
        with self.assertRaisesRegex(transfer.TransferError, "Every source"):
            transfer.prepare_manifest(self.groups, inspection, "Customer", "Merged")

    def test_password_is_required_for_discovery_and_transfer_review(self):
        with self.assertRaisesRegex(transfer.TransferError, "Developer Password"):
            transfer.inspect_sources(self.files, {**self.config, "developer_password": "wrong"})

    def test_failed_backup_restore_blocks_transfer(self):
        self.stage()
        with patch("operation_transfer._run_gbak", side_effect=transfer.TransferError("backup restore failed")):
            with self.assertRaisesRegex(transfer.TransferError, "backup restore failed"):
                self.backup()
        self.assertFalse(any(self.operation_rows().values()))

    def test_commit_connection_loss_reports_unknown_and_recheck_detects_success(self):
        self.stage()
        receipt = self.backup()
        real_connect = transfer._connect
        class LostConnection:
            def __init__(self, underlying):
                self.underlying = underlying
            def commit(self):
                self.underlying.commit()
                raise OSError("lost acknowledgement")
            def __getattr__(self, name):
                return getattr(self.underlying, name)
        with patch("operation_transfer._connect", side_effect=lambda config, write=False: LostConnection(real_connect(config, write=write))):
            with self.assertRaises(transfer.CommitOutcomeUnknown):
                self.apply(receipt)
        self.assertTrue(transfer.check_target(self.manifest_path, 102, self.config)["already_applied"])

    def test_busy_database_fails_without_waiting_or_writing(self):
        self.stage()
        receipt = self.backup()
        other = self.connect()
        try:
            cursor = other.cursor()
            cursor.execute("UPDATE CUT_LIST_DETAIL SET LENGTE=LENGTE WHERE QUOTE_NR=102")
            with self.assertRaises(fdb.DatabaseError):
                self.apply(receipt)
        finally:
            other.rollback()
            other.close()
        self.assertFalse(any(self.operation_rows().values()))

    def test_expired_backup_is_rejected(self):
        self.stage()
        receipt = self.backup()
        payload = json.loads(Path(receipt).read_text())["payload"]
        payload["verified_at"] = (datetime.now(timezone.utc) - timedelta(minutes=16)).isoformat()
        transfer._write_signed(receipt, payload)
        with self.assertRaisesRegex(transfer.TransferError, "expired"):
            self.apply(receipt)
        self.assertFalse(any(self.operation_rows().values()))

    def test_dimension_mode_and_settings_must_match_sources(self):
        self.stage()
        self.sql("UPDATE QUOTE SET SIZEMODE=0 WHERE QUOTE_NR=102")
        with self.assertRaisesRegex(transfer.TransferError, "dimension mode"):
            transfer.check_target(self.manifest_path, 102, self.config)
        self.sql("UPDATE QUOTE SET SIZEMODE=1 WHERE QUOTE_NR=102")
        self.sql("UPDATE CUTLIST SET BUILD_TYPE='2' WHERE QUOTE_NR=102")
        with self.assertRaisesRegex(transfer.TransferError, "cutting settings"):
            transfer.check_target(self.manifest_path, 102, self.config)
        self.assertFalse(any(self.operation_rows().values()))

    def test_changed_piece_ids_or_customer_cannot_be_used(self):
        self.stage()
        self.sql("INSERT INTO CUSTOMER VALUES (2,'Unrelated customer')")
        self.sql("UPDATE QUOTE SET CUSTOMER_ID=2 WHERE QUOTE_NR=102")
        with self.assertRaisesRegex(transfer.TransferError, "customer/job"):
            transfer.check_target(self.manifest_path, 102, self.config)
        self.sql("UPDATE QUOTE SET CUSTOMER_ID=1 WHERE QUOTE_NR=102")
        self.sql("UPDATE CUT_LIST_DETAIL SET UNIT_ID=1 WHERE QUOTE_NR=102")
        with self.assertRaisesRegex(transfer.TransferError, "repeated"):
            transfer.check_target(self.manifest_path, 102, self.config)

    def test_import_build_default_and_renamed_job_require_explicit_review(self):
        self.sql("UPDATE CUTLIST SET BUILD_TYPE='2' WHERE QUOTE_NR=100")
        self.sql("UPDATE CUTLIST SET BUILD_TYPE='2' WHERE QUOTE_NR=101")
        self.stage()
        self.sql("UPDATE QUOTE SET JOB_NAME='TEST - MERGE' WHERE QUOTE_NR=102")
        self.plan = transfer.check_target(self.manifest_path, 102, self.config)
        self.assertEqual(self.plan["review_differences"], [
            {"field": "Build type", "cutlist": 1, "source": "2", "target": "3"},
            {"field": "Job name", "source": "Merged Job", "target": "TEST - MERGE"}])
        receipt = self.backup()
        with self.assertRaisesRegex(transfer.TransferError, "explicitly accept"):
            self.apply(receipt)
        self.assertFalse(any(self.operation_rows().values()))
        self.assertEqual(self.apply(receipt, reviewed=True), "committed")
        self.assertEqual(self.sql("SELECT JOB_NAME FROM QUOTE WHERE QUOTE_NR=102"), [("TEST - MERGE",)])
        self.assertEqual(self.sql("SELECT QUOTE_NR,BUILD_TYPE FROM CUTLIST ORDER BY QUOTE_NR"), [(100,"2"),(101,"2"),(102,"3")])

    def test_import_blank_build_default_is_reviewable_but_mode_changes_are_blocked(self):
        self.sql("UPDATE CUTLIST SET BUILD_TYPE=NULL WHERE QUOTE_NR IN (100,101)")
        self.stage()
        self.assertEqual(self.plan["review_differences"], [
            {"field": "Build type", "cutlist": 1, "source": None, "target": "3"}])
        self.sql("UPDATE CUTLIST SET OPTIMIZE_MODE=1 WHERE QUOTE_NR=102")
        with self.assertRaisesRegex(transfer.TransferError, "optimization mode differs"):
            transfer.check_target(self.manifest_path, 102, self.config)

    def test_review_approval_cannot_be_reused_for_a_changed_target(self):
        self.stage()
        self.sql("UPDATE QUOTE SET JOB_NAME='Test A' WHERE QUOTE_NR=102")
        self.plan = transfer.check_target(self.manifest_path, 102, self.config)
        receipt = self.backup()
        self.sql("UPDATE QUOTE SET JOB_NAME='Test B' WHERE QUOTE_NR=102")
        with self.assertRaisesRegex(transfer.TransferError, "changed since review"):
            self.apply(receipt, reviewed=True)
        self.assertFalse(any(self.operation_rows().values()))

    def test_custom_triggers_and_unrecognized_schema_are_blocked(self):
        self.sql("UPDATE DB_INFO SET VERSION=25")
        with self.assertRaisesRegex(transfer.TransferError, "schema 24"):
            transfer.inspect_sources(self.files, self.config)
        self.sql("UPDATE DB_INFO SET VERSION=24")
        self.sql("CREATE TRIGGER CUSTOM_TRIGGER FOR CROSSCUTS ACTIVE BEFORE INSERT AS BEGIN END")
        with self.assertRaisesRegex(transfer.TransferError, "Custom database triggers"):
            transfer.inspect_sources(self.files, self.config)

    def test_backup_must_match_the_reviewed_plan(self):
        self.stage()
        receipt = self.backup()
        self.plan["inserts"][0]["values"][1] = 999
        with self.assertRaisesRegex(transfer.TransferError, "not verified"):
            self.apply(receipt)
        self.assertFalse(any(self.operation_rows().values()))

    def test_audit_write_failure_after_commit_is_reported_as_committed(self):
        self.stage()
        receipt = self.backup()
        real_write = transfer._write_signed
        def write(path, payload):
            if payload.get("status") == "committed":
                raise OSError("disk full")
            return real_write(path, payload)
        with patch("operation_transfer._write_signed", side_effect=write):
            self.assertEqual(self.apply(receipt), "committed_audit_pending")
        self.assertTrue(transfer.check_target(self.manifest_path, 102, self.config)["already_applied"])


if __name__ == "__main__":
    unittest.main()
