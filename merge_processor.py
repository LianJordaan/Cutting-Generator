"""Read, group and merge Cut Manager version 7 Excel exports."""

from copy import copy
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
import os
from pathlib import Path
import re
import tempfile

from openpyxl import Workbook, load_workbook


HEADERS = ("Unit ID", "L", "W", "Qty", "Edge L", "Edge W", "Holes L",
           "Holes W", "Edging Category", "Edging Name", "Remarks")
CONTROL_CELLS = ("A6", "C6", "E6", "G6", "I6", "K6")


def text(value):
    return "" if value is None else str(value).strip()


@dataclass
class SourceCutlist:
    path: Path
    sheet: object
    rows: tuple[int, ...]
    quantity: int

    @property
    def board(self):
        return text(self.sheet["E4"].value)

    @property
    def job(self):
        return text(self.sheet["C3"].value)

    @property
    def group_key(self):
        return (text(self.sheet["C4"].value).casefold(), self.board.casefold(),
                Decimal(text(self.sheet["J4"].value)),
                tuple(text(self.sheet[c].value).casefold() for c in CONTROL_CELLS))


@dataclass
class MergeFile:
    path: Path
    cutlists: list[SourceCutlist]

    @property
    def job(self):
        return self.cutlists[0].job


@dataclass
class MergeGroup:
    sources: list[SourceCutlist]

    @property
    def board(self):
        return self.sources[0].board

    @property
    def row_count(self):
        return sum(len(s.rows) for s in self.sources)

    @property
    def quantity(self):
        return sum(s.quantity for s in self.sources)


def _invalid(path, sheet, row, message):
    return ValueError(f"{path.name} / {sheet.title}, row {row}: {message}")


def _number(value, *, positive=False, integer=False, maximum=None):
    try:
        if isinstance(value, bool) or not text(value):
            raise InvalidOperation
        result = Decimal(text(value))
        if not result.is_finite() or result < 0 or (positive and result == 0):
            raise InvalidOperation
        if integer and result != result.to_integral_value():
            raise InvalidOperation
        if maximum is not None and result > maximum:
            raise InvalidOperation
        return result
    except (InvalidOperation, ValueError):
        raise ValueError(f"invalid number {value!r}") from None


def read_merge_file(file_path):
    """Load an immutable-in-practice snapshot; never edit the input workbook."""
    path = Path(file_path).resolve()
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"{path.name}: please use a Cut Manager .xlsx export (File Version 7).")
    workbook = load_workbook(path)
    try:
        cutlists = []
        for sheet in workbook:
            if (not re.fullmatch(r"CutList\d+", sheet.title, re.IGNORECASE)
                    or text(sheet["A1"].value).casefold() != "cut manager cutting list"):
                raise ValueError(f"{path.name}: {sheet.title!r} is not a Cut Manager cutlist. Use a Cut Manager .xlsx export.")
            if text(sheet["K5"].value).casefold() != "file version 7":
                raise _invalid(path, sheet, 5, "only File Version 7 is supported.")
            if tuple(text(c.value).casefold() for c in sheet[7][:11]) != tuple(h.casefold() for h in HEADERS):
                raise _invalid(path, sheet, 7, "unexpected cutlist column headings.")
            if not text(sheet["C4"].value) or not text(sheet["E4"].value):
                raise _invalid(path, sheet, 4, "board category and name are required.")
            try:
                _number(sheet["J4"].value, positive=True)
            except ValueError as exc:
                raise _invalid(path, sheet, 4, f"board thickness: {exc}") from exc
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f":
                        raise _invalid(path, sheet, cell.row, "formulas cannot be merged safely; export cell values first.")
                    if cell.column > 11 and text(cell.value):
                        raise _invalid(path, sheet, cell.row, "unexpected data beyond the Cut Manager columns.")
            if any(m.max_row >= 8 for m in sheet.merged_cells.ranges):
                raise _invalid(path, sheet, 8, "merged cells in the piece table are unsupported.")
            rows, quantity, ended = [], 0, False
            for row in sheet.iter_rows(min_row=8, max_col=11):
                values = [c.value for c in row]
                if not any(text(v) for v in values):
                    continue
                if ended:
                    raise _invalid(path, sheet, row[0].row, "unexpected data after EndOfFile.")
                if text(values[0]).casefold() == "endoffile":
                    if any(text(v) for v in values[1:]):
                        raise _invalid(path, sheet, row[0].row, "unexpected data on EndOfFile row.")
                    ended = True
                    continue
                try:
                    _number(values[0], integer=True)
                    _number(values[1], positive=True)
                    _number(values[2], positive=True)
                    count = _number(values[3], positive=True, integer=True)
                    for col in range(4, 8):
                        _number(values[col], integer=True, maximum=2 if col < 6 else None)
                    for col, header in [(8, "C5"), (9, "E5")]:
                        if (_number(values[4]) or _number(values[5])) and not text(values[col]) and not text(sheet[header].value):
                            raise ValueError("edged piece is missing its edging category/name")
                except ValueError as exc:
                    raise _invalid(path, sheet, row[0].row, str(exc)) from exc
                rows.append(row[0].row)
                quantity += int(count)
            if not rows:
                raise _invalid(path, sheet, 8, "cutlist contains no pieces.")
            cutlists.append(SourceCutlist(path, sheet, tuple(rows), quantity))
        if not cutlists:
            raise ValueError(f"{path.name}: workbook contains no cutlists.")
        return MergeFile(path, cutlists)
    finally:
        workbook.close()


def group_cutlists(files):
    """Combine matching boards in first-appearance order, retaining source row order."""
    groups = {}
    for file in files:
        for source in file.cutlists:
            if source.group_key not in groups:
                groups[source.group_key] = MergeGroup([])
            groups[source.group_key].sources.append(source)
    return list(groups.values())


def move_items(items, selected, direction):
    """Move selected rows one position without changing their relative order."""
    result = list(items)
    selected = set(selected)
    indexes = range(len(result)) if direction < 0 else range(len(result) - 1, -1, -1)
    for index in indexes:
        target = index + direction
        if index in selected and 0 <= target < len(result) and target not in selected:
            result[index], result[target] = result[target], result[index]
            selected.remove(index)
            selected.add(target)
    return result, sorted(selected)


def _copy_style(source, target):
    # Style IDs belong to a workbook; register their actual components in the new one.
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)
    target.number_format = source.number_format


def _copy_cell(source, target, style_source=None):
    target.value = source.value
    target.data_type = source.data_type
    _copy_style(style_source if style_source is not None else source, target)
    if source.comment:
        target.comment = copy(source.comment)


def _literal(cell, value):
    cell.value = value
    cell.data_type = "s"


def merge_cutlists(groups, output_path, *, customer, job, contact="", overwrite=False):
    """Write the chosen group order atomically. Caller enforces licence/dev access."""
    if not groups or not text(customer) or not text(job):
        raise ValueError("Choose cutlists and enter a customer and job name before saving.")
    output_path = Path(output_path).resolve()
    if output_path.suffix.lower() != ".xlsx":
        raise ValueError("The merged workbook must have an .xlsx extension.")
    if any(output_path == s.path for g in groups for s in g.sources):
        raise ValueError("Choose a new output filename; source files cannot be overwritten.")
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.loaded_theme = groups[0].sources[0].sheet.parent.loaded_theme
    try:
        for index, group in enumerate(groups, 1):
            original = group.sources[0].sheet
            sheet = workbook.create_sheet(f"CutList{index}")
            for row in original.iter_rows(min_row=1, max_row=7, max_col=11):
                for cell in row:
                    _copy_cell(cell, sheet.cell(cell.row, cell.column))
            for merged in original.merged_cells.ranges:
                sheet.merge_cells(str(merged))
            for key, dimension in original.column_dimensions.items():
                sheet.column_dimensions[key] = copy(dimension)
                sheet.column_dimensions[key].parent = sheet
            for row_number in range(1, 8):
                sheet.row_dimensions[row_number] = copy(original.row_dimensions[row_number])
                if original.row_dimensions[row_number].has_style:
                    _copy_style(original.row_dimensions[row_number], sheet.row_dimensions[row_number])
            sheet.sheet_format = copy(original.sheet_format)
            sheet.sheet_properties = copy(original.sheet_properties)
            sheet.page_setup = copy(original.page_setup)
            sheet.page_margins = copy(original.page_margins)
            sheet.print_options = copy(original.print_options)
            sheet.freeze_panes = "B8"
            for address, value in [("C2", text(customer)), ("C3", text(job)), ("J2", text(contact))]:
                _literal(sheet[address], value)
            destination_row = 8
            for source in group.sources:
                for source_row in source.rows:
                    for column in range(1, 12):
                        cell = source.sheet.cell(source_row, column)
                        style = cell if cell.has_style else source.sheet.cell(8, column)
                        _copy_cell(cell, sheet.cell(destination_row, column), style)
                    # Resolve header defaults before combining different edging types.
                    if _number(sheet.cell(destination_row, 5).value) or _number(sheet.cell(destination_row, 6).value):
                        for column, header in [(9, "C5"), (10, "E5")]:
                            cell = sheet.cell(destination_row, column)
                            if not text(cell.value):
                                _literal(cell, text(source.sheet[header].value))
                    sheet.row_dimensions[destination_row].height = source.sheet.row_dimensions[source_row].height
                    destination_row += 1
            sheet.cell(destination_row, 1, "EndOfFile")
            sheet.print_area = f"A1:K{destination_row}"
            sheet.print_title_rows = "1:7"
        temporary = None
        try:
            with tempfile.NamedTemporaryFile(dir=output_path.parent, suffix=".xlsx", delete=False) as handle:
                temporary = Path(handle.name)
            workbook.save(temporary)
            if overwrite:
                os.replace(temporary, output_path)
            else:
                # Atomic exclusive creation; a concurrent save cannot overwrite an existing file.
                os.link(temporary, output_path)
                temporary.unlink()
            temporary = None
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)
        return output_path
    finally:
        workbook.close()
